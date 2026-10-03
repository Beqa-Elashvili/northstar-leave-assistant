"""Text extraction from the supplied DOCX and PDF documents (read-only).

Both formats produce the same structure: an ordered list of `Block`s (title, heading,
paragraph, list item, table) plus the document's metadata table (code, version, effective
date, ...). Structure is taken from the documents' own formatting:

- DOCX: paragraph styles (Title, Heading 1/2, List Bullet, Normal) and tables, in body order.
  DOCX has no reliable page numbers, so `page` is always None for DOCX blocks.
- PDF (text-based, no OCR): font size/weight per line (title 19pt bold, article 14pt bold,
  sub-article 11.4pt bold, body 11pt), PyMuPDF table detection, running header/footer
  removed. `page` is the 1-based physical page, which equals the printed "გვერდი N".
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import pymupdf
from docx import Document as open_docx
from docx.table import Table
from docx.text.paragraph import Paragraph

# PyMuPDF otherwise print()s a package recommendation to stdout, which would corrupt
# stdio-based protocols and clutter the CLI.
if hasattr(pymupdf, "no_recommend_layout"):
    pymupdf.no_recommend_layout()

BlockKind = Literal["title", "heading", "paragraph", "list_item", "table"]

# "4.4 Title", "4. Title", "10.5 Title" (policies) and "ა.2 Title" / "ა. Title" (FAQ).
_SECTION_RE = re.compile(r"^\s*((?:\d+(?:\.\d+)*)|(?:[ა-ჰ](?:\.\d+)?))\.?\s+(.+?)\s*$")
_METADATA_KEYS = {
    "კომპანია": "company", "დოკუმენტის კოდი": "doc_code", "ვერსია": "version", "ძალაშია": "effective",
    "ბოლო განახლება": "last_updated", "მფლობელი": "owner", "თანამფლობელი": "co_owner",
    "სტატუსი": "status", "დამტკიცებულია": "approved_by",
}


@dataclass
class Block:
    kind: BlockKind
    text: str
    page: int | None = None
    level: int = 0                       # headings: 1 = article, 2 = sub-article
    section: str | None = None           # headings: "4.4", "ა.2"
    heading_title: str | None = None     # headings: text without the number
    rows: list[list[str]] | None = None  # tables


@dataclass
class ExtractedDocument:
    document: str
    file_type: Literal["docx", "pdf"]
    title: str
    metadata: dict[str, str]
    blocks: list[Block] = field(default_factory=list)
    page_count: int | None = None


def parse_heading(text: str) -> tuple[str | None, str]:
    match = _SECTION_RE.match(text)
    if not match:
        return None, text.strip()
    return match.group(1), match.group(2)


def render_table(rows: list[list[str]]) -> str:
    """Table as text that keeps cell/header association, e.g. 'ადგილი: ლონდონი; ლიმიტი ერთ ღამეზე: 180 ...'."""
    rows = [[_clean(c) for c in r] for r in rows if any(_clean(c) for c in r)]
    if not rows:
        return ""
    if len(rows[0]) == 2 and not _looks_like_header(rows[0]):
        return "\n".join(f"{k}: {v}" for k, v in rows)
    header, body = rows[0], rows[1:]
    lines = [" | ".join(header)]
    for r in body:
        lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, r) if v))
    return "\n".join(lines)


def _looks_like_header(row: list[str]) -> bool:
    # Data tables in these documents have a short header row (e.g. "ადგილი | ლიმიტი ერთ ღამეზე");
    # metadata tables start with "კომპანია".
    return row[0] not in _METADATA_KEYS


def _clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _metadata_from_rows(rows: list[list[str]]) -> dict[str, str]:
    meta = {}
    for r in rows:
        if len(r) >= 2 and _clean(r[0]) in _METADATA_KEYS:
            meta[_METADATA_KEYS[_clean(r[0])]] = _clean(r[1])
    return meta


# --- DOCX ------------------------------------------------------------------------------------------

def extract_docx(path: Path) -> ExtractedDocument:
    doc = open_docx(str(path))
    blocks: list[Block] = []
    metadata: dict[str, str] = {}
    title = ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = _clean(para.text)
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style == "Title":
                title = title or text
                blocks.append(Block("title", text))
            elif style.startswith("Heading"):
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                section, heading_title = parse_heading(text)
                blocks.append(Block("heading", text, level=level, section=section, heading_title=heading_title))
            elif style.startswith("List"):
                blocks.append(Block("list_item", text))
            else:
                blocks.append(Block("paragraph", text))
        elif tag == "tbl":
            table = Table(child, doc)
            rows = [[_clean(c.text) for c in row.cells] for row in table.rows]
            if not metadata and _metadata_from_rows(rows):
                metadata = _metadata_from_rows(rows)
            blocks.append(Block("table", render_table(rows), rows=rows))
    return ExtractedDocument(path.name, "docx", title, metadata, blocks)


# --- PDF -------------------------------------------------------------------------------------------

_HEADER_MAX_Y = 45.0
_FOOTER_MIN_Y = 795.0
_RUNNING_TEXT_MAX_SIZE = 8.5


@dataclass
class _Line:
    text: str
    size: float
    bold: bool
    y0: float
    y1: float


def _line_kind(line: _Line, page_no: int) -> tuple[BlockKind, int]:
    if line.bold and line.size >= 17 and page_no == 1:
        return "title", 0
    if line.bold and line.size >= 13:
        return "heading", 1
    if line.bold and line.size >= 11.2:
        return "heading", 2
    return "paragraph", 0


def _in_rect(line: _Line, rect) -> bool:
    return rect.y0 - 1 <= line.y0 and line.y1 <= rect.y1 + 1


def extract_pdf(path: Path) -> ExtractedDocument:
    pdf = pymupdf.open(str(path))
    blocks: list[Block] = []
    metadata: dict[str, str] = {}
    title_parts: list[str] = []
    try:
        for page_no, page in enumerate(pdf, start=1):
            tables = page.find_tables().tables
            table_rects = [pymupdf.Rect(t.bbox) for t in tables]
            items: list[tuple[float, object]] = [(r.y0, t) for r, t in zip(table_rects, tables)]

            for raw_block in page.get_text("dict")["blocks"]:
                for raw_line in raw_block.get("lines", []):
                    spans = [s for s in raw_line["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    line = _Line(
                        text=_clean("".join(s["text"] for s in raw_line["spans"])),
                        size=spans[0]["size"], bold=bool(spans[0]["flags"] & 16) or "Bold" in spans[0]["font"],
                        y0=raw_line["bbox"][1], y1=raw_line["bbox"][3],
                    )
                    running = line.size <= _RUNNING_TEXT_MAX_SIZE and (
                        line.y0 < _HEADER_MAX_Y or line.y0 > _FOOTER_MIN_Y)
                    if running or any(_in_rect(line, r) for r in table_rects):
                        continue
                    items.append((line.y0, line))

            items.sort(key=lambda item: item[0])
            current: Block | None = None
            last_line: _Line | None = None
            for _, item in items:
                if not isinstance(item, _Line):  # a table
                    rows = [[_clean(c) for c in row] for row in item.extract()]
                    current, last_line = None, None
                    if not metadata and page_no == 1 and _metadata_from_rows(rows):
                        metadata = _metadata_from_rows(rows)
                    _append_table(blocks, rows, page_no)
                    continue
                kind, level = _line_kind(item, page_no)
                if kind == "title":
                    title_parts.append(item.text)
                    current, last_line = None, item
                    continue
                gap = item.y0 - last_line.y1 if last_line else 99.0
                same_style = last_line is not None and current is not None and \
                    (item.bold, round(item.size)) == (last_line.bold, round(last_line.size))
                # Wrapped lines of the same paragraph are ~4.8pt apart; new paragraphs/list items ~11pt.
                if same_style and gap < item.size * 0.6:
                    current.text = f"{current.text} {item.text}"
                    if kind == "heading":
                        current.section, current.heading_title = parse_heading(current.text)
                else:
                    if kind == "heading":
                        section, heading_title = parse_heading(item.text)
                        current = Block("heading", item.text, page=page_no, level=level, section=section,
                                        heading_title=heading_title)
                    else:
                        current = Block("paragraph", item.text, page=page_no)
                    blocks.append(current)
                last_line = item
        page_count = pdf.page_count
    finally:
        pdf.close()
    title = " ".join(title_parts)
    if title:
        blocks.insert(0, Block("title", title, page=1))
    return ExtractedDocument(path.name, "pdf", title, metadata, blocks, page_count)


def _append_table(blocks: list[Block], rows: list[list[str]], page_no: int) -> None:
    """A table that continues on the next page repeats its header row: merge it into the previous table."""
    previous = blocks[-1] if blocks else None
    if previous is not None and previous.kind == "table" and previous.rows and rows \
            and previous.rows[0] == rows[0] and (previous.page or 0) == page_no - 1:
        previous.rows.extend(rows[1:])
        previous.text = render_table(previous.rows)
        return
    blocks.append(Block("table", render_table(rows), page=page_no, rows=rows))


# --- entry point ---------------------------------------------------------------------------------

def extract_document(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    raise ValueError(f"unsupported document type: {path.name}")


def list_documents(directory: Path) -> list[Path]:
    return sorted(p for p in directory.iterdir()
                  if p.suffix.lower() in (".docx", ".pdf") and not p.name.startswith("~$"))
